"""Unit tests for deterministic vendor file generators."""

import io
import json
import time
from datetime import date

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from fakes.storage import FakeStorage
from omni_retail.generators.vendor_files.generator import (
    generate_payload,
    make_filename,
)
from omni_retail.ingestion.common.paths import BUCKET_LANDING
from omni_retail.ingestion.files.schemas import (
    HISTORICAL_ORDERS,
    PARTNER_PRODUCTS,
    SUPPLIER_PRICES,
    SUPPLIER_STOCK,
)
from omni_retail.ingestion.files.validation import validate_payload

RUN_DATE = date(2026, 9, 11)


def test_same_seed_produces_identical_payload() -> None:
    first = generate_payload(SUPPLIER_PRICES, seed=7, rows=50, run_date=RUN_DATE)
    second = generate_payload(SUPPLIER_PRICES, seed=7, rows=50, run_date=RUN_DATE)

    assert first == second


def test_different_seed_produces_different_payload() -> None:
    first = generate_payload(SUPPLIER_PRICES, seed=7, rows=50, run_date=RUN_DATE)
    second = generate_payload(SUPPLIER_PRICES, seed=8, rows=50, run_date=RUN_DATE)

    assert first != second


def test_generated_csv_passes_validation() -> None:
    payload = generate_payload(SUPPLIER_PRICES, seed=7, rows=100, run_date=RUN_DATE)

    result = validate_payload(SUPPLIER_PRICES, payload)

    assert result.file_valid, result.file_errors
    assert result.row_count == 100
    assert result.bad_rows == ()


def test_generated_json_passes_validation() -> None:
    payload = generate_payload(PARTNER_PRODUCTS, seed=7, rows=100, run_date=RUN_DATE)

    result = validate_payload(PARTNER_PRODUCTS, payload)

    assert result.file_valid, result.file_errors
    assert result.row_count == 100
    assert result.bad_rows == ()


def test_generated_csv_rows_are_parseable_and_sorted_by_schema() -> None:
    payload = generate_payload(SUPPLIER_PRICES, seed=3, rows=10, run_date=RUN_DATE)
    header = payload.decode().splitlines()[0]

    assert header == "supplier_id,sku,price,currency,valid_from"


def test_generated_json_is_array_of_schema_records() -> None:
    payload = generate_payload(PARTNER_PRODUCTS, seed=3, rows=10, run_date=RUN_DATE)
    document = json.loads(payload)

    assert isinstance(document, list) and len(document) == 10
    assert set(document[0]) == {field.name for field in PARTNER_PRODUCTS.fields}


def test_generated_parquet_passes_validation() -> None:
    payload = generate_payload(HISTORICAL_ORDERS, seed=7, rows=100, run_date=RUN_DATE)

    result = validate_payload(HISTORICAL_ORDERS, payload)

    assert result.file_valid, result.file_errors
    assert result.row_count == 100
    assert result.bad_rows == ()


def test_generated_parquet_has_typed_columns() -> None:
    payload = generate_payload(HISTORICAL_ORDERS, seed=7, rows=10, run_date=RUN_DATE)

    table = pq.read_table(io.BytesIO(payload))

    assert table.column_names == ["order_id", "customer_id", "status", "order_total", "order_date"]
    assert table.schema.field("order_id").type == pa.int64()
    assert table.schema.field("order_total").type == pa.decimal128(12, 2)
    assert table.schema.field("order_date").type == pa.date32()


def test_parquet_payload_is_deterministic() -> None:
    first = generate_payload(HISTORICAL_ORDERS, seed=11, rows=25, run_date=RUN_DATE)
    second = generate_payload(HISTORICAL_ORDERS, seed=11, rows=25, run_date=RUN_DATE)

    assert first == second


def test_generated_xlsx_passes_validation() -> None:
    payload = generate_payload(SUPPLIER_STOCK, seed=7, rows=100, run_date=RUN_DATE)

    result = validate_payload(SUPPLIER_STOCK, payload)

    assert result.file_valid, result.file_errors
    assert result.row_count == 100
    assert result.bad_rows == ()


def test_generated_xlsx_has_header_and_typed_cells() -> None:
    payload = generate_payload(SUPPLIER_STOCK, seed=3, rows=5, run_date=RUN_DATE)

    workbook = load_workbook(io.BytesIO(payload), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))

    assert rows[0] == ("supplier_id", "sku", "quantity", "updated_at")
    assert isinstance(rows[1][2], int)
    assert rows[1][3] is not None


def test_xlsx_payload_is_deterministic_across_time() -> None:
    # openpyxl stamps the workbook and zip entries with the current time;
    # the generator normalizes them, and zip timestamps have 2s resolution,
    # so the sleep must exceed it to prove the normalization works.
    first = generate_payload(SUPPLIER_STOCK, seed=11, rows=5, run_date=RUN_DATE)
    time.sleep(2.5)
    second = generate_payload(SUPPLIER_STOCK, seed=11, rows=5, run_date=RUN_DATE)

    assert first == second


def test_filename_is_deterministic() -> None:
    assert make_filename(SUPPLIER_PRICES, seed=7, rows=50, run_date=RUN_DATE) == (
        "supplier-prices_20260911_s7_n50.csv"
    )
    assert make_filename(PARTNER_PRODUCTS, seed=7, rows=50, run_date=RUN_DATE) == (
        "partner-products_20260911_s7_n50.json"
    )
    assert make_filename(HISTORICAL_ORDERS, seed=7, rows=50, run_date=RUN_DATE) == (
        "historical-orders_20260911_s7_n50.parquet"
    )
    assert make_filename(SUPPLIER_STOCK, seed=7, rows=50, run_date=RUN_DATE) == (
        "supplier-stock_20260911_s7_n50.xlsx"
    )


def test_upload_writes_incoming_object() -> None:
    storage = FakeStorage()
    payload = generate_payload(SUPPLIER_PRICES, seed=7, rows=5, run_date=RUN_DATE)

    from omni_retail.generators.vendor_files.generator import upload_payload

    upload_payload(storage, SUPPLIER_PRICES, payload, seed=7, rows=5, run_date=RUN_DATE)

    key = "supplier-prices/incoming/supplier-prices_20260911_s7_n5.csv"
    assert storage.stored_objects()[(BUCKET_LANDING, key)] == payload


def test_rows_must_be_positive() -> None:
    with pytest.raises(ValueError, match="rows"):
        generate_payload(SUPPLIER_PRICES, seed=7, rows=0, run_date=RUN_DATE)
